"""Acceptance tests for the offline company-demo DOCX cleaner."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path
from uuid import uuid4

import pytest
from docx import Document

from evaluation.company_demo.text_cleaning import CleaningValidationError, clean_company_demo


def _write_docx(path: Path) -> None:
    """Create a local DOCX fixture with paragraphs, placeholders, and a table."""
    document = Document()
    document.add_paragraph("  公司   管理制度  ")
    document.add_paragraph("适用单位：××公司；日期：____年____月____日")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "流程"
    table.cell(0, 1).text = "负责人"
    table.cell(1, 0).text = "审批"
    table.cell(1, 1).text = "部门经理"
    document.save(path)


def _write_manifest(root: Path, source_path: Path) -> None:
    """Create the minimal manifest consumed by the cleaner."""
    document_id = str(uuid4())
    relative_path = source_path.relative_to(root).as_posix()
    manifest = {
        "schema_version": 1,
        "corpus_id": "test-company-demo",
        "selection": {"departments": ["human_resources"]},
        "documents": [
            {
                "source_document_id": document_id,
                "department": "human_resources",
                "document_type": "policy",
                "title": "测试制度",
                "source_filename": source_path.name,
                "repository_path": relative_path,
                "source_format": "docx",
                "authority": "unverified_reference",
                "status": "needs_review",
                "sensitivity": "internal_demo",
                "sha256": hashlib.sha256(source_path.read_bytes()).hexdigest(),
            }
        ],
    }
    (root / "corpus_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")


def test_clean_company_demo_preserves_structure_and_reports_placeholders(tmp_path: Path) -> None:
    """Given a DOCX fixture, cleaning keeps tables and placeholders in normalized text."""
    source = tmp_path / "documents" / "human_resources" / "test.docx"
    source.parent.mkdir(parents=True)
    _write_docx(source)
    _write_manifest(tmp_path, source)
    before_hash = hashlib.sha256(source.read_bytes()).hexdigest()

    report = clean_company_demo(tmp_path)

    normalized = tmp_path / "normalized" / "human_resources" / "test.txt"
    text = normalized.read_text(encoding="utf-8")
    assert "公司 管理制度" in text
    assert "适用单位：××公司；日期：____年____月____日" in text
    assert "表格行: 流程 | 负责人" in text
    assert "表格行: 审批 | 部门经理" in text
    assert report.document_count == 1
    assert report.placeholder_hits == 2
    assert hashlib.sha256(source.read_bytes()).hexdigest() == before_hash

    report_json = json.loads((tmp_path / "cleaning_report.json").read_text(encoding="utf-8"))
    assert report_json["documents"][0]["status"] == "succeeded"
    assert report_json["documents"][0]["input_sha256"] == before_hash


def test_clean_company_demo_rejects_invalid_docx(tmp_path: Path) -> None:
    """Given a corrupt DOCX, cleaning fails instead of emitting a misleading text file."""
    source = tmp_path / "documents" / "human_resources" / "broken.docx"
    source.parent.mkdir(parents=True)
    source.write_bytes(b"not a docx")
    _write_manifest(tmp_path, source)

    with pytest.raises(CleaningValidationError, match="failed"):
        clean_company_demo(tmp_path)
    assert not (tmp_path / "normalized" / "human_resources" / "broken.txt").exists()
