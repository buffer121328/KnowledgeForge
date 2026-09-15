"""Acceptance tests for code-based company-demo department expansion."""

from __future__ import annotations

import base64
import hashlib
import json
from pathlib import Path
from uuid import NAMESPACE_URL, uuid5

from docx import Document
from evaluation.company_demo.department_expansion import (
    inspect_docx_objects,
    prepare_company_demo_department,
)

_TINY_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUB"
    "AScY42YAAAAASUVORK5CYII="
)


def _write_docx(path: Path, *, with_image: bool) -> None:
    document = Document()
    document.add_paragraph("行政管理制度")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "事项"
    table.cell(0, 1).text = "负责人"
    if with_image:
        image = path.with_suffix(".png")
        image.write_bytes(_TINY_PNG)
        document.add_picture(str(image))
    document.save(path)


def _selection(path: Path, names: list[str]) -> None:
    candidates = []
    for name in names:
        candidates.append(
            {
                "source_document_id": str(
                    uuid5(NAMESPACE_URL, f"knowledgeforge/company-demo/administration/{name}")
                ),
                "source_filename": name,
                "document_type": "policy",
                "title": Path(name).stem,
                "authority": "unverified_reference",
                "status": "needs_review",
                "sensitivity": "internal_demo",
            }
        )
    path.write_text(
        json.dumps(
            {
                "schema_version": 1,
                "department": "administration",
                "department_label": "行政管理部",
                "candidates": candidates,
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )


def _empty_corpus(root: Path) -> None:
    source = root / "documents" / "human_resources" / "existing.docx"
    source.parent.mkdir(parents=True)
    _write_docx(source, with_image=False)
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    (root / "corpus_manifest.json").write_text(
        json.dumps(
            {
                "schema_version": 1,
                "corpus_id": "test-company-demo",
                "selection": {
                    "departments": ["human_resources"],
                    "department_labels": {"human_resources": "人力资源部"},
                },
                "documents": [
                    {
                        "source_document_id": str(uuid5(NAMESPACE_URL, "existing")),
                        "department": "human_resources",
                        "department_label": "人力资源部",
                        "document_type": "policy",
                        "title": "existing",
                        "source_filename": source.name,
                        "repository_path": "documents/human_resources/existing.docx",
                        "source_format": "docx",
                        "authority": "unverified_reference",
                        "status": "needs_review",
                        "sensitivity": "internal_demo",
                        "sha256": digest,
                    }
                ],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )


def test_docx_inspection_detects_images_and_accepts_plain_tables(tmp_path: Path) -> None:
    plain = tmp_path / "plain.docx"
    image = tmp_path / "image.docx"
    _write_docx(plain, with_image=False)
    _write_docx(image, with_image=True)

    assert inspect_docx_objects(plain).has_image_content is False
    inspected = inspect_docx_objects(image)
    assert inspected.has_image_content is True
    assert inspected.media_count == 1
    assert inspected.drawing_count >= 1


def test_preparation_skips_image_word_and_only_manifests_plain_docx(tmp_path: Path) -> None:
    source_root = tmp_path / "source"
    source_root.mkdir()
    plain = source_root / "plain.docx"
    image = source_root / "image.docx"
    _write_docx(plain, with_image=False)
    _write_docx(image, with_image=True)
    corpus_root = tmp_path / "corpus"
    _empty_corpus(corpus_root)
    selection = tmp_path / "selection.json"
    _selection(selection, [plain.name, image.name])

    report = prepare_company_demo_department(
        source_root,
        corpus_root,
        selection_path=selection,
    )

    assert report.included_count == 1
    assert report.skipped_count == 1
    statuses = {record.source_filename: record.status for record in report.records}
    assert statuses == {
        "plain.docx": "included",
        "image.docx": "skipped_image_content",
    }
    assert (corpus_root / "documents/administration/plain.docx").is_file()
    assert not (corpus_root / "documents/administration/image.docx").exists()
    assert (corpus_root / "normalized/administration/plain.txt").is_file()
    assert not (corpus_root / "normalized/administration/image.txt").exists()

    manifest = json.loads((corpus_root / "corpus_manifest.json").read_text())
    administration = [
        item for item in manifest["documents"] if item["department"] == "administration"
    ]
    assert [item["source_filename"] for item in administration] == ["plain.docx"]
    preparation = json.loads(
        (corpus_root / "administration_preparation_report.json").read_text()
    )
    skipped = next(
        item
        for item in preparation["candidates"]
        if item["status"] == "skipped_image_content"
    )
    assert skipped["reason_code"] == "docx_contains_image_or_drawing"
    assert skipped["media_count"] == 1
